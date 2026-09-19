import os
import io
import csv
from collections import Counter
from fastapi import APIRouter, Response
from fastapi.responses import FileResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

from database import get_db

router = APIRouter(tags=["Reports"])

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

@router.get("/api/reports/kpi")
async def get_kpi():
    async with get_db() as conn:
        c = await conn.cursor()
        await c.execute("SELECT COUNT(*) FROM audits")
        total = (await c.fetchone())[0]

        await c.execute("SELECT COUNT(*) FROM audits WHERE slm_verdict = 'NOK'")
        incidents = (await c.fetchone())[0]

        compliant = total - incidents
        rate = round((compliant / total * 100), 1) if total > 0 else 100.0

        await c.execute("SELECT line, timestamp, slm_analysis FROM audits ORDER BY id DESC LIMIT 10")
        recent = [{"line": r[0], "timestamp": r[1], "slm_analysis": r[2]} for r in await c.fetchall()]
    
    return {
        "total_audits": total,
        "compliance_rate": rate,
        "incidents_count": incidents,
        "recent_audits": recent
    }

@router.get("/api/audits/export/excel")
async def export_excel():
    async with get_db() as conn:
        c = await conn.cursor()
        await c.execute("SELECT * FROM audits ORDER BY id DESC")
        rows = [dict(r) for r in await c.fetchall()]

    wb = openpyxl.Workbook()
    
    # --- ARKUSZ 1: DASHBOARD  ---
    ws_dash = wb.active
    ws_dash.title = "Dashboard Admin"
    ws_dash.sheet_view.showGridLines = False
    
    # Nagłówek Dashboardu
    ws_dash.merge_cells('A1:J2')
    title_cell = ws_dash['A1']
    title_cell.value = "QUALITY AUDIT ENTERPRISE - DASHBOARD Admin IFS FOOD v8"
    title_cell.font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Obliczenia 
    total_audits = len(rows)
    nok_audits = sum(1 for r in rows if r.get("slm_verdict", "") == "NOK")
    ok_audits = total_audits - nok_audits
    compliance = round((ok_audits / total_audits * 100), 1) if total_audits > 0 else 100.0
    
    kpi_data = [
        ("Suma przeprowadzonych audytów:", total_audits),
        ("Wskaźnik zgodności IFS (Compliance):", f"{compliance}%"),
        ("Audyty w pełni zgodne (OK):", ok_audits),
        ("Złamania KO / Niezgodności (NOK):", nok_audits)
    ]
    
    for idx, (label, val) in enumerate(kpi_data, start=4):
        ws_dash.merge_cells(f'B{idx}:D{idx}')
        ws_dash[f'B{idx}'] = label
        ws_dash[f'B{idx}'].font = Font(size=12, bold=True, color="334155")
        ws_dash[f'B{idx}'].alignment = Alignment(horizontal="right")
        
        ws_dash[f'E{idx}'] = val
        ws_dash[f'E{idx}'].font = Font(size=14, bold=True, color="16A34A" if idx != 7 else "DC2626")
        ws_dash[f'E{idx}'].alignment = Alignment(horizontal="left")

    lines_counter = Counter([r.get("line", "Nieznana") for r in rows])
    
    ws_dash['A30'] = "Linia"
    ws_dash['B30'] = "Liczba Audytów"
    row_idx = 31
    for line, count in lines_counter.items():
        ws_dash[f'A{row_idx}'] = line
        ws_dash[f'B{row_idx}'] = count
        row_idx += 1
        
    ws_dash['D30'] = "Status"
    ws_dash['E30'] = "Ilość"
    ws_dash['D31'] = "ZGODNE (OK)"
    ws_dash['E31'] = ok_audits
    ws_dash['D32'] = "NIEZGODNE (NOK)"
    ws_dash['E32'] = nok_audits

    # Wykres Słupkowy
    if lines_counter:
        bar_chart = BarChart()
        bar_chart.type = "col"
        bar_chart.style = 10
        bar_chart.title = "Liczba Inspekcji na Linie"
        bar_chart.width = 14
        bar_chart.height = 7
        data1 = Reference(ws_dash, min_col=2, min_row=30, max_row=row_idx-1)
        cats1 = Reference(ws_dash, min_col=1, min_row=31, max_row=row_idx-1)
        bar_chart.add_data(data1, titles_from_data=True)
        bar_chart.set_categories(cats1)
        ws_dash.add_chart(bar_chart, "B10")
    
    # Wykres Kołowy
    if total_audits > 0:
        pie_chart = PieChart()
        pie_chart.title = "Podział Zgodności"
        pie_chart.width = 10
        pie_chart.height = 7
        labels = Reference(ws_dash, min_col=4, min_row=31, max_row=32)
        data2 = Reference(ws_dash, min_col=5, min_row=30, max_row=32)
        pie_chart.add_data(data2, titles_from_data=True)
        pie_chart.set_categories(labels)
        ws_dash.add_chart(pie_chart, "G10")

    # --- ARKUSZ 2: SZCZEGÓŁOWE DANE ---
    ws_data = wb.create_sheet(title="Rejestr Danych Audytowych")
    
    headers = [
        "ID", "Data i Czas", "Audytor", "Linia", "Zmiana", "Strefa", "Zdrowie PR15",
        "CCP1 Fe", "CCP1 Non-Fe", "CCP1 SS", "Ramię Odrzutu", "Kosz Zamknięty", 
        "Czystość (GMP)", "Orzeczenie AI", "Status Zapisu"
    ]
    ws_data.append(headers)
    
    for r in rows:
        ws_data.append([
            r.get("id", ""), r.get("timestamp", ""), r.get("auditor_id", ""), r.get("line", ""), 
            r.get("shift", ""), r.get("zone", ""), r.get("health_ok", ""), r.get("ccp1_fe_ok", ""), 
            r.get("ccp1_nonfe_ok", ""), r.get("ccp1_ss_ok", ""), r.get("ccp1_reject_ok", ""), 
            r.get("ccp1_bin_locked", ""), r.get("gmp_cleanliness_ok", ""), r.get("slm_analysis", ""), 
            r.get("record_status", "ZABLOKOWANY")
        ])
        
    if total_audits > 0:
        tab = Table(displayName="TabelaAudytow", ref=f"A1:O{total_audits+1}")
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        tab.tableStyleInfo = style
        ws_data.add_table(tab)
    
    for col_idx in range(1, len(headers)+1):
        ws_data.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 22

    out_file = os.path.join(BASE_DIR, "Raport_Dashboard_IFS_Enterprise.xlsx")
    wb.save(out_file)
    
    return FileResponse(
        out_file, 
        filename="Raport_Dashboard_IFS_Enterprise.xlsx", 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@router.get("/api/audits/export/powerbi/csv")
async def export_powerbi_csv():
    """Generuje zoptymalizowany pod kątem Microsoft Power BI plik CSV (kodowanie UTF-8-SIG z BOM)"""
    async with get_db() as conn:
        c = await conn.cursor()
        await c.execute("""
            SELECT id, timestamp, auditor_id, line, shift, zone,
                   health_ok, dispense_no, line_status,
                   ccp1_fe_ok, ccp1_nonfe_ok, ccp1_ss_ok, ccp1_reject_ok, ccp1_bin_locked,
                   ccp2_magnet_ok, ccp3_sieve_ok,
                   gmp_cleanliness_ok, gmp_wood_score, gmp_foreign_score, gmp_waste_ok,
                   bhp_estop_ok, bhp_atex_ok, bhp_hot_cip_ok, bhp_evac_ppoz_ok, bhp_status,
                   slm_verdict, risk_level, record_status, compliance_verdict,
                   total_score_pct, audit_score, ko_failed, notes
            FROM audits ORDER BY id DESC
        """)
        rows = [dict(r) for r in await c.fetchall()]

    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    
    headers = [
        "Audit_ID", "Timestamp", "Data", "Godzina", "Rok", "Miesiac", "Kwartal",
        "Audytor", "Linia_Produkcyjna", "Zmiana", "Strefa", "Status_Linii",
        "Werdykt_SLM", "Poziom_Ryzyka", "Status_Decyzji", "Zgodnosc_Standardu",
        "CCP_Zgodny", "GMP_Zgodny", "BHP_Zgodny", "Blokada_KO",
        "Wynik_Punktowy", "Uwagi"
    ]
    writer.writerow(headers)

    for r in rows:
        ts = r.get("timestamp") or ""
        date_part = ts.split(" ")[0] if " " in ts else ts
        time_part = ts.split(" ")[1] if " " in ts else ""
        year_part = date_part.split("-")[0] if "-" in date_part else ""
        month_part = date_part.split("-")[1] if "-" in date_part else ""
        
        q_part = ""
        if month_part.isdigit():
            m_int = int(month_part)
            q_part = f"Q{(m_int - 1) // 3 + 1}"

        ccp_ok = "TAK" if (r.get("ccp1_fe_ok") == "OK" and r.get("ccp1_reject_ok") == "OK" and r.get("ccp2_magnet_ok") == "OK") else "NIE"
        gmp_ok = "TAK" if r.get("gmp_cleanliness_ok") == "OK" else "NIE"
        bhp_ok = "TAK" if r.get("bhp_status") == "OK" else "NIE"
        ko_flag = "TAK" if r.get("ko_failed") == 1 else "NIE"

        writer.writerow([
            r.get("id"),
            ts,
            date_part,
            time_part,
            year_part,
            month_part,
            q_part,
            r.get("auditor_id") or "Nieprzypisany",
            r.get("line") or "Nieznana",
            r.get("shift") or "1",
            r.get("zone") or "Standard",
            r.get("line_status") or "PRODUKCJA",
            r.get("slm_verdict") or "OK",
            r.get("risk_level") or "NISKIE",
            r.get("compliance_verdict") or "ZATWIERDZONY",
            r.get("record_status") or "ZABLOKOWANY",
            ccp_ok,
            gmp_ok,
            bhp_ok,
            ko_flag,
            r.get("total_score_pct") or 100,
            (r.get("notes") or "").replace("\n", " ").replace("\r", "")
        ])

    csv_data = output.getvalue()
    encoded_bytes = csv_data.encode("utf-8-sig")

    return Response(
        content=encoded_bytes,
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": "attachment; filename=Raport_Audyty_IFS_PowerBI.csv",
            "Access-Control-Allow-Origin": "*"
        }
    )

@router.get("/api/audits/export/powerbi/feed")
async def export_powerbi_feed():
    """Zwraca tabelaryczny JSON zoptymalizowany pod kątem Power BI Web Connector"""
    async with get_db() as conn:
        c = await conn.cursor()
        await c.execute("""
            SELECT id, timestamp, auditor_id, line, shift, zone,
                   slm_verdict, risk_level, record_status, compliance_verdict,
                   total_score_pct, ko_failed, notes
            FROM audits ORDER BY id DESC
        """)
        rows = [dict(r) for r in await c.fetchall()]

    for r in rows:
        ts = r.get("timestamp") or ""
        date_part = ts.split(" ")[0] if " " in ts else ts
        r["date"] = date_part
        r["year"] = date_part.split("-")[0] if "-" in date_part else ""
        r["month"] = date_part.split("-")[1] if "-" in date_part else ""

    return {"value": rows}

